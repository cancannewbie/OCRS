"""Generated-only, immutable XLSX snapshots of human-confirmed orders.

The application supplies a new system-controlled generation filename. Existing
files are never silently overwritten: identical bytes are an idempotent retry;
different bytes, including manual edits, require a new generation filename.
"""

from __future__ import annotations

import errno
import hashlib
import os
import re
import stat
import tempfile
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation, localcontext
from pathlib import Path
from typing import Any, TypedDict, cast
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ocrs.domain import CURRENCY_MINOR_UNITS

SCHEMA_VERSION = "1"
LOCK_TIMEOUT = 2.0
_FIXED_TIME = datetime(2000, 1, 1)
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_ORDER_HEADERS = (
    "order_id",
    "version",
    "status",
    "customer",
    "external_id",
    "currency",
    "occurred_at",
    "created_at",
    "updated_at",
    "items_total",
    "items_total_exact",
)
_ITEM_HEADERS = (
    "order_id",
    "line_id",
    "version",
    "status",
    "sku",
    "name",
    "quantity",
    "quantity_exact",
    "unit",
    "unit_price",
    "unit_price_exact",
    "line_total",
    "line_total_exact",
)


class ExportError(ValueError):
    """Safe error code with an explicit retryability hint."""

    def __init__(self, code: str, message: str, *, retryable: bool = False) -> None:
        self.code = code
        self.retryable = retryable
        super().__init__(message)


class ExportMetadata(TypedDict):
    filename: str
    sha256: str
    order_count: int
    item_count: int
    schema_version: str


def _decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, str | int | Decimal):
        raise ExportError("export_invalid_amount", "Amounts must be exact decimal values.")
    try:
        result = Decimal(value)
    except InvalidOperation as exc:
        raise ExportError("export_invalid_amount", "An amount is invalid.") from exc
    if not result.is_finite():
        raise ExportError("export_invalid_amount", "Amounts must be finite.")
    parts = result.as_tuple()
    if (
        len(parts.digits) > 64
        or result.adjusted() > 64
        or result.adjusted() < -64
        or not isinstance(parts.exponent, int)
        or abs(parts.exponent) > 64
        or result < 0
    ):
        raise ExportError("export_invalid_amount", "An amount exceeds supported precision.")
    return result


def _exact(value: Decimal | None) -> str | None:
    return None if value is None else format(value, "f")


def _numeric(value: Decimal | None) -> Decimal | None:
    # Excel guarantees only 15 significant decimal digits. Always retain the
    # authoritative exact string next to the optional convenience numeric cell.
    if value is None:
        return None
    digits = value.as_tuple().digits
    while digits and digits[-1] == 0:
        digits = digits[:-1]
    return value if len(digits) <= 15 and -99 <= value.adjusted() <= 99 else None


def _text(value: object) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise ExportError("export_invalid_record", "Text fields must contain strings.")
    escaped = _BAD_XML.sub(lambda match: f"\\u{ord(match.group()):04x}", value)
    if len(escaped) > 32_767:
        raise ExportError("export_text_too_long", "A text field exceeds the Excel cell limit.")
    return escaped


def _set_value(cell: Cell, value: str | int | Decimal | None) -> None:
    cell.value = value
    if isinstance(value, str):
        # This includes every user/model string, not just a blacklist of formula
        # prefixes. Tabs, newlines, spaces, '=','+','-','@' remain literal text.
        cell.data_type = "s"
        cell.number_format = "@"
    elif isinstance(value, Decimal):
        precision = min(12, max(2, -int(value.as_tuple().exponent)))
        cell.number_format = "#,##0." + "0" * precision


def _append(sheet: Worksheet, values: list[str | int | Decimal | None]) -> None:
    row = sheet.max_row + 1
    for column, value in enumerate(values, 1):
        _set_value(cast(Cell, sheet.cell(row, column)), value)


def _prepare_sheet(sheet: Worksheet, headers: tuple[str, ...]) -> None:
    for column, header in enumerate(headers, 1):
        cell = cast(Cell, sheet.cell(1, column))
        _set_value(cell, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[cell.column_letter].width = min(36, max(18, len(header) + 4))
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _identifier(record: dict[str, Any], key: str, fallback: str | None = None) -> str:
    value = record.get(key, record.get(fallback) if fallback else None)
    if not isinstance(value, str) or not value:
        raise ExportError("export_invalid_record", "Stable record identifiers are required.")
    result = _text(value)
    assert result is not None
    return result


def _build_workbook(orders: list[dict[str, Any]]) -> tuple[Workbook, int]:
    if not isinstance(orders, list) or not all(isinstance(order, dict) for order in orders):
        raise ExportError("export_invalid_record", "An order list is required.")
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    workbook.properties.creator = "OCRS"
    workbook.properties.title = "OCRS confirmed order snapshot"
    workbook.properties.description = (
        "Generated-only immutable snapshot. Copy before manual edits. "
        "The database is authoritative; *_exact columns retain exact decimals. "
        f"Schema version {SCHEMA_VERSION}."
    )
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    order_sheet = workbook.create_sheet("orders")
    item_sheet = workbook.create_sheet("items")
    _prepare_sheet(order_sheet, _ORDER_HEADERS)
    _prepare_sheet(item_sheet, _ITEM_HEADERS)
    seen_orders: set[str] = set()
    item_count = 0
    ordered = sorted(orders, key=lambda row: _identifier(row, "order_id", "id"))
    for order in ordered:
        order_id = _identifier(order, "order_id", "id")
        if order_id in seen_orders:
            raise ExportError("export_duplicate_record", "An order identifier is duplicated.")
        seen_orders.add(order_id)
        status = order.get("status")
        version = order.get("version")
        if not isinstance(status, str) or status not in {"confirmed", "cancelled"}:
            raise ExportError("export_unconfirmed_record", "Only formal orders can be exported.")
        currency = order.get("currency")
        if not isinstance(currency, str) or currency not in CURRENCY_MINOR_UNITS:
            raise ExportError("export_invalid_currency", "A supported currency is required.")
        if isinstance(version, bool) or not isinstance(version, int) or version < 1:
            raise ExportError("export_invalid_record", "A positive order version is required.")
        items = order.get("items")
        if not isinstance(items, list) or not all(isinstance(item, dict) for item in items):
            raise ExportError("export_invalid_record", "An order must contain an item list.")
        seen_lines: set[str] = set()
        total: Decimal | None = Decimal(0)
        # Inputs are bounded to 64 digits and +/-64 scale. This context covers
        # products plus all possible rows in an Excel worksheet without rounding.
        with localcontext() as context:
            context.prec = 256
            for item in sorted(items, key=lambda row: _identifier(row, "line_id")):
                line_id = _identifier(item, "line_id")
                if line_id in seen_lines:
                    raise ExportError("export_duplicate_record", "A line identifier is duplicated.")
                seen_lines.add(line_id)
                quantity = _decimal(item.get("quantity"))
                price = _decimal(item.get("unit_price"))
                quantum = Decimal(1).scaleb(-CURRENCY_MINOR_UNITS[currency])
                line_total = (
                    None
                    if quantity is None or price is None
                    else (quantity * price).quantize(quantum, rounding=ROUND_HALF_UP)
                )
                total = None if total is None or line_total is None else total + line_total
                _append(
                    item_sheet,
                    [
                        order_id,
                        line_id,
                        version,
                        str(status),
                        _text(item.get("sku")),
                        _text(item.get("name")),
                        _numeric(quantity),
                        _exact(quantity),
                        _text(item.get("unit")),
                        _numeric(price),
                        _exact(price),
                        _numeric(line_total),
                        _exact(line_total),
                    ],
                )
                item_count += 1
                if item_count >= 1_048_576:
                    raise ExportError("export_too_many_rows", "The snapshot exceeds Excel limits.")
        _append(
            order_sheet,
            [
                order_id,
                version,
                str(status),
                _text(order.get("customer")),
                _text(order.get("external_id")),
                _text(order.get("currency")),
                _text(order.get("occurred_at")),
                _text(order.get("created_at")),
                _text(order.get("updated_at")),
                _numeric(total),
                _exact(total),
            ],
        )
    if len(ordered) >= 1_048_576:
        raise ExportError("export_too_many_rows", "The snapshot exceeds Excel limits.")
    for sheet in (order_sheet, item_sheet):
        sheet.auto_filter.ref = sheet.dimensions
    return workbook, item_count


def _canonicalize(source: Path, destination: Path) -> None:
    """Stabilize package timestamps so identical snapshot retries compare bytes."""
    with ZipFile(source) as original, ZipFile(destination, "w", ZIP_DEFLATED) as normalized:
        for name in sorted(original.namelist()):
            content = original.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                    content,
                )
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            normalized.writestr(info, content)


def _verify(path: Path, order_count: int, item_count: int) -> None:
    with ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ExportError("export_verification_failed", "The workbook package is invalid.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if workbook.sheetnames != ["orders", "items"]:
            raise ExportError("export_verification_failed", "Workbook sheets are incomplete.")
        for name, headers, count in (
            ("orders", _ORDER_HEADERS, order_count),
            ("items", _ITEM_HEADERS, item_count),
        ):
            sheet = workbook[name]
            if sheet.max_row != count + 1 or tuple(cell.value for cell in sheet[1]) != headers:
                raise ExportError("export_verification_failed", "Workbook rows are incomplete.")
            for row in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise ExportError("export_verification_failed", "Formula cells are forbidden.")
            numeric_pairs = ((10, 11),) if name == "orders" else ((7, 8), (10, 11), (12, 13))
            for row in sheet.iter_rows(min_row=2):
                for numeric_column, exact_column in numeric_pairs:
                    numeric, exact = row[numeric_column - 1].value, row[exact_column - 1].value
                    if numeric is not None and Decimal(str(numeric)) != Decimal(str(exact)):
                        raise ExportError(
                            "export_verification_failed", "A numeric value lost precision."
                        )
    finally:
        workbook.close()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _existing_snapshot(destination: Path, digest: str) -> bool:
    try:
        mode = destination.lstat().st_mode
    except FileNotFoundError:
        return False
    if stat.S_ISREG(mode) and _sha256(destination) == digest:
        return True
    raise ExportError("export_target_exists", "Use a new generation filename; a file exists.")


def _publish(source: Path, destination: Path, digest: str) -> None:
    # os.replace has a check-then-overwrite race with manual/non-cooperating
    # writers. A same-filesystem hard link atomically publishes the complete
    # immutable generation only if the destination is absent. No unsafe fallback.
    try:
        os.link(source, destination)
    except FileExistsError:
        if not _existing_snapshot(destination, digest):
            raise ExportError(
                "export_locked", "The export target changed; retry later.", retryable=True
            ) from None
    except OSError as exc:
        if exc.errno in {errno.EXDEV, errno.ENOSYS, errno.ENOTSUP, errno.EOPNOTSUPP} or getattr(
            exc, "winerror", None
        ) in {1, 50}:
            raise ExportError(
                "export_publish_unsupported", "This filesystem cannot publish safe snapshots."
            ) from exc
        raise


def write_workbook(orders: list[dict[str, Any]], destination: Path) -> ExportMetadata:
    """Validate, lock, verify and atomically publish an immutable local snapshot.

    Only the application may choose destination, inside its private export root.
    Advisory file locking supports cooperating processes on one host; this is not
    a distributed/network-filesystem writer. A locked file is a retryable error.
    """
    destination = Path(destination)
    if destination.suffix.lower() != ".xlsx" or destination.is_symlink():
        raise ExportError(
            "export_invalid_destination", "A controlled XLSX destination is required."
        )
    try:
        if not destination.parent.is_dir() or destination.parent.is_symlink():
            raise ExportError("export_invalid_destination", "The export directory is unavailable.")
        with FileLock(str(destination) + ".lock", timeout=LOCK_TIMEOUT):
            return _write_locked(orders, destination)
    except Timeout as exc:
        raise ExportError(
            "export_locked", "The export target is busy; retry later.", retryable=True
        ) from exc
    except PermissionError as exc:
        raise ExportError(
            "export_locked", "The export target is unavailable; retry later.", retryable=True
        ) from exc
    except (OSError, BadZipFile) as exc:
        raise ExportError(
            "export_io_error", "The snapshot could not be written.", retryable=True
        ) from exc
    except (ValueError, TypeError, KeyError, IndexError, OverflowError, SyntaxError) as exc:
        if isinstance(exc, ExportError):
            raise
        raise ExportError("export_invalid_record", "The snapshot contains invalid data.") from exc


def _write_locked(orders: list[dict[str, Any]], destination: Path) -> ExportMetadata:
    workbook, item_count = _build_workbook(orders)
    temporary: list[Path] = []
    try:
        for _ in range(2):
            fd, name = tempfile.mkstemp(
                prefix=".ocrs-export-", suffix=".xlsx", dir=destination.parent
            )
            os.close(fd)
            temporary.append(Path(name))
        raw, final = temporary
        workbook.save(raw)
        _canonicalize(raw, final)
        _verify(final, len(orders), item_count)
        with final.open("rb") as stream:
            os.fsync(stream.fileno())
        digest = _sha256(final)
        metadata: ExportMetadata = {
            "filename": destination.name,
            "sha256": digest,
            "order_count": len(orders),
            "item_count": item_count,
            "schema_version": SCHEMA_VERSION,
        }
        if _existing_snapshot(destination, digest):
            return metadata
        _publish(final, destination, digest)
        if os.name == "posix":
            directory_fd = os.open(destination.parent, os.O_RDONLY | os.O_DIRECTORY)
            try:
                os.fsync(directory_fd)
            finally:
                os.close(directory_fd)
        return metadata
    finally:
        workbook.close()
        for path in temporary:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass  # Cleanup must not hide the original failure or undo a published file.
