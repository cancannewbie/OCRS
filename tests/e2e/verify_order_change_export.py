"""Inspect a real browser download after a synthetic create/amend/cancel workflow."""

import argparse
from pathlib import Path

from openpyxl import load_workbook


def verify_export(
    path: Path, order_id: str, line_id: str, customer: str, external_id: str, item_name: str
) -> None:
    """Scope assertions by unique IDs; other viewport tests share the demo service."""
    workbook = load_workbook(path, data_only=False)
    try:
        assert {"orders", "items"}.issubset(workbook.sheetnames)
        orders = workbook["orders"]
        items = workbook["items"]
        order_headers = [cell.value for cell in orders[1]]
        item_headers = [cell.value for cell in items[1]]
        order_rows = [
            dict(zip(order_headers, row, strict=True)) for row in orders.iter_rows(min_row=2)
        ]
        item_rows = [
            dict(zip(item_headers, row, strict=True)) for row in items.iter_rows(min_row=2)
        ]
        matches = [row for row in order_rows if row["order_id"].value == order_id]
        assert len(matches) == 1, "Amending and cancelling must not append another order."
        assert sum(row["customer"].value == customer for row in order_rows) == 1
        order = matches[0]
        assert order["version"].value == 3
        assert order["status"].value == "cancelled"
        assert order["customer"].value == customer
        assert order["external_id"].value == external_id
        assert order["external_id"].data_type == "s", "Leading-zero references must stay text."
        assert order["currency"].value == "CNY"
        assert order["items_total_exact"].value == "38.75"
        assert order["items_total_exact"].data_type == "s"

        lines = [row for row in item_rows if row["order_id"].value == order_id]
        assert len(lines) == 1, "The amended line must appear once, including after cancellation."
        assert sum(row["line_id"].value == line_id for row in item_rows) == 1
        line = lines[0]
        assert line["line_id"].value == line_id
        assert line["version"].value == 3
        assert line["status"].value == "cancelled"
        assert line["sku"].value == "DEMO-001"
        assert line["name"].value == item_name
        assert line["name"].data_type == "s", "Formula-like names must remain literal text."
        assert line["unit"].value == "box"
        for column, expected in (
            ("quantity_exact", "3.125"),
            ("unit_price_exact", "12.40"),
            ("line_total_exact", "38.75"),
        ):
            assert line[column].value == expected, f"{column} must retain exact decimal text."
            assert line[column].data_type == "s"
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path)
    parser.add_argument("order_id")
    parser.add_argument("line_id")
    parser.add_argument("customer")
    parser.add_argument("external_id")
    parser.add_argument("item_name")
    arguments = parser.parse_args()
    verify_export(
        arguments.path,
        arguments.order_id,
        arguments.line_id,
        arguments.customer,
        arguments.external_id,
        arguments.item_name,
    )


if __name__ == "__main__":
    main()
