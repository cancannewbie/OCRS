"""Draw an invented evidence image and inspect the actual browser-downloaded XLSX."""

import argparse
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


def make_image(destination: Path) -> None:
    """No customer screenshots, network resources, or system fonts are used."""
    image = Image.new("RGB", (760, 1060), "#eef2f5")
    draw = ImageDraw.Draw(image)
    title = ImageFont.load_default(size=32)
    text = ImageFont.load_default(size=24)
    small = ImageFont.load_default(size=19)
    draw.rectangle((0, 0, 760, 142), fill="#174f48")
    draw.text((38, 36), "SYNTHETIC TEST DATA", font=title, fill="white")
    draw.text(
        (38, 87), "Invented evidence. No real customer information.", font=small, fill="white"
    )
    draw.rounded_rectangle((32, 182, 698, 586), radius=22, fill="white")
    draw.text((58, 211), "Fictional customer: Sample Alpha", font=text, fill="#24332f")
    for index, line in enumerate(
        [
            "Product: Fictional sample box",
            "SKU: DEMO-001",
            "Quantity: 2.5 boxes",
            "Unit price: CNY 12.40",
            "External reference: 000042",
            "This image is generated for browser tests.",
        ]
    ):
        draw.text((58, 273 + index * 44), line, font=text, fill="#3f514b")
    draw.rounded_rectangle((124, 636, 728, 852), radius=22, fill="#d8eade")
    for index, line in enumerate(
        [
            "Human review is still required.",
            "The demo provider does not read pixels.",
            "Its candidates are fictional examples.",
            "No model API is called by this test.",
        ]
    ):
        draw.text((151, 670 + index * 42), line, font=small, fill="#234f41")
    draw.text((38, 968), "OCRS / DISPOSABLE OFFLINE ACCEPTANCE FIXTURE", font=small, fill="#57665f")
    destination.parent.mkdir(parents=True, exist_ok=True)
    image.save(destination, format="PNG")


def verify_export(path: Path, customer: str, external_id: str, item_name: str) -> None:
    """Check semantic workbook contents, not just a download filename or ZIP magic."""
    workbook = load_workbook(path, data_only=False)
    try:
        assert {"orders", "items"}.issubset(workbook.sheetnames)
        orders = workbook["orders"]
        items = workbook["items"]
        order_headers = [cell.value for cell in orders[1]]
        item_headers = [cell.value for cell in items[1]]
        order_rows = [
            dict(zip(order_headers, row, strict=True)) for row in orders.iter_rows(min_row=2)
        ]
        item_rows = [
            dict(zip(item_headers, row, strict=True)) for row in items.iter_rows(min_row=2)
        ]
        matches = [row for row in order_rows if row["customer"].value == customer]
        assert len(matches) == 1, "The reviewed order must appear exactly once."
        order = matches[0]
        assert order["status"].value == "confirmed"
        assert order["version"].value == 1
        assert order["external_id"].value == external_id
        assert order["external_id"].data_type == "s", "Leading zeros must remain text."
        assert order["currency"].value == "CNY"
        assert order["items_total_exact"].value == "31.00"
        assert all(
            row["customer"].value
            not in {"DEMO ONLY — fictional customer", "SYNTHETIC Vision Buyer"}
            for row in order_rows
        ), "Unconfirmed demo or network-fixture candidates must not leak into the workbook."
        lines = [row for row in item_rows if row["order_id"].value == order["order_id"].value]
        assert len(lines) == 1
        line = lines[0]
        assert line["sku"].value == "DEMO-001"
        assert line["name"].value == item_name
        assert line["name"].data_type == "s", "Formula-like item names must remain literal text."
        assert line["quantity_exact"].value == "2.5"
        assert line["unit"].value == "box"
        assert line["unit_price_exact"].value == "12.40"
        assert line["line_total_exact"].value == "31.00"
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    actions = parser.add_subparsers(dest="action", required=True)
    image = actions.add_parser("image")
    image.add_argument("path", type=Path)
    verify = actions.add_parser("verify-export")
    verify.add_argument("path", type=Path)
    verify.add_argument("customer")
    verify.add_argument("external_id")
    verify.add_argument("item_name")
    arguments = parser.parse_args()
    if arguments.action == "image":
        make_image(arguments.path)
    else:
        verify_export(
            arguments.path, arguments.customer, arguments.external_id, arguments.item_name
        )


if __name__ == "__main__":
    main()
