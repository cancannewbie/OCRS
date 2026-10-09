"""Synthetic-only capture and generated snapshot regression tests."""

from __future__ import annotations

import errno
import io
import multiprocessing
import os
from decimal import Decimal
from pathlib import Path
from typing import Any
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from filelock import FileLock
from openpyxl import load_workbook
from PIL import Image

from ocrs.capture import CaptureError, InboxWatcher, validate_image
from ocrs.exporter import ExportError, write_workbook


def image_bytes(format_name: str = "PNG") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (12, 9), "blue").save(buffer, format=format_name)
    return buffer.getvalue()


def formal_order(**changes: Any) -> dict[str, Any]:
    order: dict[str, Any] = {
        "id": "order-fiction-1",
        "version": 1,
        "status": "confirmed",
        "customer": "Fictional Customer",
        "external_id": "000042",
        "currency": "CNY",
        "occurred_at": None,
        "created_at": "2026-01-01T00:00:00+00:00",
        "updated_at": "2026-01-01T00:00:00+00:00",
        "items": [
            {
                "line_id": "line-1",
                "sku": "001",
                "name": "Fictional item",
                "quantity": "3",
                "unit": "box",
                "unit_price": "0.10",
            },
            {
                "line_id": "line-2",
                "sku": None,
                "name": "Fictional item B",
                "quantity": "1",
                "unit": "box",
                "unit_price": "0.20",
            },
        ],
    }
    order.update(changes)
    return order


@pytest.mark.parametrize(
    ("format_name", "mime", "extension"),
    [("PNG", "image/png", ".png"), ("JPEG", "image/jpeg", ".jpg"), ("WEBP", "image/webp", ".webp")],
)
def test_image_validation_reads_truthful_format(
    format_name: str, mime: str, extension: str
) -> None:
    data = image_bytes(format_name)
    info = validate_image(data)
    assert (info.mime, info.extension, info.width, info.height) == (mime, extension, 12, 9)
    assert len(info.sha256) == 64
    assert info.sha256 == validate_image(data).sha256


@pytest.mark.parametrize("data", [b"", b"not an image", b"\x89PNG\r\n\x1a\n"])
def test_invalid_images_fail_safely(data: bytes) -> None:
    with pytest.raises(CaptureError, match="image|nonempty") as error:
        validate_image(data)
    assert error.value.code == "image_invalid"


def test_byte_pixel_and_format_limits() -> None:
    for data, options, code in (
        (image_bytes(), {"max_bytes": 10}, "image_too_large"),
        (image_bytes(), {"max_pixels": 100}, "image_too_many_pixels"),
        (image_bytes("GIF"), {}, "image_type_unsupported"),
    ):
        with pytest.raises(CaptureError) as error:
            validate_image(data, **options)
        assert error.value.code == code


def test_decompression_bomb_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 50)
    with pytest.raises(CaptureError) as error:
        validate_image(image_bytes())
    assert error.value.code == "image_too_many_pixels"


def test_animated_png_is_rejected() -> None:
    buffer = io.BytesIO()
    first = Image.new("RGB", (3, 3), "blue")
    first.save(buffer, "PNG", save_all=True, append_images=[Image.new("RGB", (3, 3), "red")])
    with pytest.raises(CaptureError) as error:
        validate_image(buffer.getvalue())
    assert error.value.code == "image_animated"


def test_truncated_compressed_image_is_rejected() -> None:
    with pytest.raises(CaptureError) as error:
        validate_image(image_bytes("JPEG")[:-20])
    assert error.value.code == "image_invalid"


def test_watcher_settles_changed_files_and_deduplicates(tmp_path: Path) -> None:
    imported: list[tuple[bytes, str, str]] = []
    watcher = InboxWatcher(tmp_path, lambda *args: imported.append(args))
    image = tmp_path / "drop.JPG"
    image.write_bytes(b"partial")
    assert watcher.scan_once()["pending"] == 1
    image.write_bytes(image_bytes())
    assert watcher.scan_once()["pending"] == 1
    assert watcher.scan_once()["imported"] == 1
    assert imported == [(image_bytes(), "image/png", "inbox:drop.JPG")]
    assert watcher.scan_once()["imported"] == 0
    assert len(imported) == 1


def test_watcher_reports_invalid_once_and_recovers_after_change(tmp_path: Path) -> None:
    path = tmp_path / "broken.png"
    path.write_bytes(b"invalid")
    watcher = InboxWatcher(tmp_path, lambda *_: None)
    assert watcher.scan_once()["pending"] == 1
    assert watcher.scan_once()["errors"] == [{"filename": "broken.png", "code": "image_invalid"}]
    assert watcher.scan_once()["errors"] == []
    path.write_bytes(image_bytes())
    watcher.scan_once()
    assert watcher.scan_once()["imported"] == 1


def test_watcher_skips_nested_directories_symlinks_and_other_files(tmp_path: Path) -> None:
    nested = tmp_path / "nested"
    nested.mkdir()
    real = nested / "hidden.png"
    real.write_bytes(image_bytes())
    (tmp_path / "alias.png").symlink_to(real)
    (tmp_path / "text.txt").write_bytes(image_bytes())
    imported: list[bytes] = []
    watcher = InboxWatcher(tmp_path, lambda data, *_: imported.append(data))
    watcher.scan_once()
    assert watcher.scan_once()["imported"] == 0
    assert not imported
    alias = tmp_path / "linked-inbox"
    alias.symlink_to(nested, target_is_directory=True)
    assert InboxWatcher(alias, lambda *_: None).scan_once()["errors"][0]["code"] == "inbox_unsafe"


def test_watcher_limits_callback_retries(tmp_path: Path) -> None:
    (tmp_path / "valid.png").write_bytes(image_bytes())
    attempts = 0

    def fail(*_: object) -> None:
        nonlocal attempts
        attempts += 1
        raise RuntimeError("Private callback exception must not leak")

    watcher = InboxWatcher(tmp_path, fail)
    for _ in range(8):
        result = watcher.scan_once()
        assert "Private" not in str(result)
    assert attempts == 3


def test_watcher_bounds_reads_and_resumes_batches(tmp_path: Path) -> None:
    for i in range(5):
        (tmp_path / f"image-{i}.png").write_bytes(image_bytes())
    imported: list[bytes] = []
    watcher = InboxWatcher(tmp_path, lambda data, *_: imported.append(data), max_files=2)
    for _ in range(7):
        assert watcher.scan_once()["imported"] <= 2
    assert len(imported) == 5
    large = tmp_path / "too-large.png"
    large.write_bytes(b"a" * 101)
    bounded = InboxWatcher(tmp_path, lambda *_: None, max_bytes=100)
    bounded.scan_once()
    result = bounded.scan_once()
    assert {"filename": "too-large.png", "code": "image_too_large"} in result["errors"]


def test_watcher_detects_file_swap_before_open(tmp_path: Path) -> None:
    path = tmp_path / "drop.png"
    path.write_bytes(image_bytes())
    imported: list[bytes] = []
    watcher = InboxWatcher(tmp_path, lambda data, *_: imported.append(data))
    watcher.scan_once()
    original_open = os.open

    def swap_open(name: Any, flags: int, *args: Any, **kwargs: Any) -> int:
        if name == "drop.png":
            path.unlink()
            path.write_bytes(image_bytes("JPEG"))
        return original_open(name, flags, *args, **kwargs)

    with patch("ocrs.capture.os.open", side_effect=swap_open):
        # Preserve supports_dir_fd membership when replacing the function object.
        with patch("ocrs.capture.os.supports_dir_fd", {os.open}):
            assert watcher.scan_once()["imported"] == 0
    assert not imported


def test_watcher_reports_unreadable_file_with_bounded_retries(tmp_path: Path) -> None:
    path = tmp_path / "drop.png"
    path.write_bytes(image_bytes())
    watcher = InboxWatcher(tmp_path, lambda *_: None)
    watcher.scan_once()
    original_open = os.open
    attempts = 0

    def denied_open(name: Any, flags: int, *args: Any, **kwargs: Any) -> int:
        nonlocal attempts
        if name == "drop.png":
            attempts += 1
            raise PermissionError("Private path")
        return original_open(name, flags, *args, **kwargs)

    with patch("ocrs.capture.os.open", side_effect=denied_open):
        with patch("ocrs.capture.os.supports_dir_fd", {os.open}):
            assert watcher.scan_once()["errors"] == [
                {"filename": "drop.png", "code": "inbox_read_failed"}
            ]
            for _ in range(7):
                watcher.scan_once()
    assert attempts == 3


def test_watcher_portable_fallback_does_not_open_directory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    (tmp_path / "valid.png").write_bytes(image_bytes())
    imported: list[bytes] = []
    watcher = InboxWatcher(tmp_path, lambda data, *_: imported.append(data))
    original_open = os.open

    def reject_directory_open(name: Any, flags: int, *args: Any, **kwargs: Any) -> int:
        assert Path(name) != tmp_path
        return original_open(name, flags, *args, **kwargs)

    monkeypatch.setattr(os, "supports_fd", set())
    monkeypatch.setattr(os, "open", reject_directory_open)
    assert watcher.scan_once()["pending"] == 1
    assert watcher.scan_once()["imported"] == 1
    assert imported == [image_bytes()]


def test_portable_fallback_rechecks_directory_after_read(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "valid.png").write_bytes(image_bytes())
    watcher = InboxWatcher(inbox, lambda *_: pytest.fail("Changed directory must not import"))
    monkeypatch.setattr(os, "supports_fd", set())
    watcher.scan_once()
    original_stat = watcher._stat_file
    stats = 0

    def swap_after_stat(fd: int | None, name: str) -> os.stat_result:
        nonlocal stats
        result = original_stat(fd, name)
        stats += 1
        if stats == 2:
            inbox.rename(tmp_path / "original-inbox")
            inbox.mkdir()
        return result

    monkeypatch.setattr(watcher, "_stat_file", swap_after_stat)
    result = watcher.scan_once()
    assert result["imported"] == 0
    assert result["errors"] == [{"filename": "", "code": "inbox_changed"}]


def sheet_rows(path: Path, name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=False)
    try:
        values = list(workbook[name].values)
        return [dict(zip(values[0], row, strict=True)) for row in values[1:]]
    finally:
        workbook.close()


def test_export_exact_decimal_totals_and_stable_keys(tmp_path: Path) -> None:
    path = tmp_path / "snapshot.xlsx"
    metadata = write_workbook([formal_order()], path)
    assert (metadata["order_count"], metadata["item_count"], metadata["schema_version"]) == (
        1,
        2,
        "1",
    )
    orders = sheet_rows(path, "orders")
    assert orders[0]["items_total_exact"] == "0.50"
    assert Decimal(str(orders[0]["items_total"])) == Decimal("0.50")
    assert orders[0]["external_id"] == "000042"
    items = sheet_rows(path, "items")
    assert items[0]["line_total_exact"] == "0.30"
    assert items[0]["order_id"] == "order-fiction-1"
    assert items[0]["line_id"] == "line-1"
    assert items[0]["version"] == 1 and items[0]["status"] == "confirmed"
    workbook = load_workbook(path)
    for sheet in workbook:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
    workbook.close()


@pytest.mark.parametrize(
    "malicious",
    [
        '=HYPERLINK("https://invalid.example","x")',
        "+1+2",
        "-1+2",
        "@SUM(1,2)",
        " \t\r\n=1+2",
        "\x00=1+2",
        "\x1b+1+2",
        "\u200b=1+2",
        "00001",
        "汉字\n=1+2",
    ],
)
def test_every_exported_text_is_literal(tmp_path: Path, malicious: str) -> None:
    path = tmp_path / "snapshot.xlsx"
    record = formal_order(customer=malicious, external_id=malicious)
    for item in record["items"]:
        item.update(name=malicious, sku=malicious, unit=malicious)
    write_workbook([record], path)
    workbook = load_workbook(path, data_only=False)
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
                if isinstance(cell.value, str):
                    assert cell.data_type == "s"
    workbook.close()
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/"):
                assert b"<f>" not in archive.read(name)


def test_export_high_precision_stays_exact_without_misleading_numeric(tmp_path: Path) -> None:
    order = formal_order()
    order["items"] = [
        {
            "line_id": "line-precise",
            "sku": None,
            "name": "Fictional precise item",
            "quantity": "12345678901234567890.123456789",
            "unit": "g",
            "unit_price": "0.01",
        }
    ]
    path = tmp_path / "snapshot.xlsx"
    write_workbook([order], path)
    item = sheet_rows(path, "items")[0]
    assert item["quantity"] is None
    assert item["line_total"] is None
    assert item["line_total_exact"] == "123456789012345678.90"


@pytest.mark.parametrize(
    ("currency", "price", "expected"),
    [
        ("CNY", "0.005", "0.02"),
        ("JPY", "1.5", "4"),
    ],
)
def test_export_rounds_each_line_before_summing(
    tmp_path: Path,
    currency: str,
    price: str,
    expected: str,
) -> None:
    order = formal_order(currency=currency)
    for item in order["items"]:
        item.update(quantity="1", unit_price=price)
    path = tmp_path / "rounded.xlsx"
    write_workbook([order], path)
    assert sheet_rows(path, "orders")[0]["items_total_exact"] == expected


@pytest.mark.parametrize("value", [0.1, float("nan"), "NaN", "Infinity", "-1", "1e1000"])
def test_export_rejects_float_and_invalid_decimal(tmp_path: Path, value: object) -> None:
    order = formal_order()
    order["items"][0]["unit_price"] = value
    with pytest.raises(ExportError) as error:
        write_workbook([order], tmp_path / "snapshot.xlsx")
    assert error.value.code == "export_invalid_amount"
    assert not (tmp_path / "snapshot.xlsx").exists()


def test_export_missing_amounts_remain_unknown(tmp_path: Path) -> None:
    order = formal_order()
    order["items"][0]["unit_price"] = None
    path = tmp_path / "snapshot.xlsx"
    write_workbook([order], path)
    assert sheet_rows(path, "orders")[0]["items_total"] is None
    assert sheet_rows(path, "orders")[0]["items_total_exact"] is None


def test_export_refuses_candidates_and_duplicate_ids(tmp_path: Path) -> None:
    with pytest.raises(ExportError) as candidate_error:
        write_workbook([formal_order(status="review_required")], tmp_path / "candidate.xlsx")
    assert candidate_error.value.code == "export_unconfirmed_record"
    with pytest.raises(ExportError) as duplicate_error:
        write_workbook([formal_order(), formal_order()], tmp_path / "duplicate.xlsx")
    assert duplicate_error.value.code == "export_duplicate_record"
    cancelled = tmp_path / "cancelled.xlsx"
    write_workbook([formal_order(status="cancelled", version=2)], cancelled)
    assert sheet_rows(cancelled, "orders")[0]["status"] == "cancelled"


def test_export_is_deterministic_and_never_overwrites_edits(tmp_path: Path) -> None:
    path = tmp_path / "snapshot.xlsx"
    record = formal_order()
    first = write_workbook([record], path)
    original = path.read_bytes()
    record["items"].reverse()
    second = write_workbook([record], path)
    assert first == second and path.read_bytes() == original
    workbook = load_workbook(path)
    workbook["orders"]["D2"] = "Manual fictional edit"
    workbook.save(path)
    workbook.close()
    edited = path.read_bytes()
    with pytest.raises(ExportError) as error:
        write_workbook([record], path)
    assert error.value.code == "export_target_exists"
    assert path.read_bytes() == edited


def test_export_lock_timeout_is_retryable(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "locked.xlsx"
    monkeypatch.setattr("ocrs.exporter.LOCK_TIMEOUT", 0.01)
    with FileLock(str(path) + ".lock"):
        with pytest.raises(ExportError) as error:
            write_workbook([formal_order()], path)
    assert error.value.code == "export_locked" and error.value.retryable
    assert write_workbook([formal_order()], path)["order_count"] == 1


def test_failed_atomic_publish_leaves_no_partial_and_can_retry(tmp_path: Path) -> None:
    path = tmp_path / "snapshot.xlsx"
    with patch("ocrs.exporter.os.link", side_effect=PermissionError("private path")):
        with pytest.raises(ExportError) as error:
            write_workbook([formal_order()], path)
    assert error.value.code == "export_locked" and "private path" not in str(error.value)
    assert not path.exists()
    assert not list(tmp_path.glob(".ocrs-export-*"))
    assert write_workbook([formal_order()], path)["item_count"] == 2


def test_failed_verification_never_publishes(tmp_path: Path) -> None:
    path = tmp_path / "snapshot.xlsx"
    with patch("ocrs.exporter._verify", side_effect=ExportError("export_verification_failed", "x")):
        with pytest.raises(ExportError) as error:
            write_workbook([formal_order()], path)
    assert error.value.code == "export_verification_failed"
    assert not path.exists() and not list(tmp_path.glob(".ocrs-export-*"))


def test_export_refuses_symlink_destination(tmp_path: Path) -> None:
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"Do not overwrite")
    link = tmp_path / "linked.xlsx"
    link.symlink_to(existing)
    with pytest.raises(ExportError) as error:
        write_workbook([formal_order()], link)
    assert error.value.code == "export_invalid_destination"
    assert existing.read_bytes() == b"Do not overwrite"


def test_atomic_publication_cannot_overwrite_a_racing_manual_file(tmp_path: Path) -> None:
    path = tmp_path / "snapshot.xlsx"
    original_link = os.link

    def racing_link(source: Path, destination: Path) -> None:
        destination.write_bytes(b"Manual file created just before publication")
        original_link(source, destination)

    with patch("ocrs.exporter.os.link", side_effect=racing_link):
        with pytest.raises(ExportError) as error:
            write_workbook([formal_order()], path)
    assert error.value.code == "export_target_exists"
    assert path.read_bytes() == b"Manual file created just before publication"
    assert not list(tmp_path.glob(".ocrs-export-*"))


def test_unsupported_atomic_publication_fails_without_unsafe_fallback(tmp_path: Path) -> None:
    path = tmp_path / "snapshot.xlsx"
    with patch("ocrs.exporter.os.link", side_effect=OSError(errno.EOPNOTSUPP, "unsupported")):
        with pytest.raises(ExportError) as error:
            write_workbook([formal_order()], path)
    assert error.value.code == "export_publish_unsupported"
    assert not error.value.retryable
    assert not path.exists()


def test_empty_snapshot_has_both_header_only_sheets(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    metadata = write_workbook([], path)
    assert metadata["order_count"] == metadata["item_count"] == 0
    assert sheet_rows(path, "orders") == sheet_rows(path, "items") == []


def _process_write(path: str, queue: Any) -> None:
    try:
        queue.put(write_workbook([formal_order()], Path(path))["sha256"])
    except ExportError as error:
        queue.put(error.code)


def test_concurrent_processes_publish_one_idempotent_snapshot(tmp_path: Path) -> None:
    context = multiprocessing.get_context("spawn")
    queue = context.Queue()
    path = tmp_path / "shared.xlsx"
    processes = [context.Process(target=_process_write, args=(str(path), queue)) for _ in range(3)]
    for process in processes:
        process.start()
    results = [queue.get(timeout=30) for _ in processes]
    for process in processes:
        process.join(timeout=30)
        assert process.exitcode == 0
    assert len(set(results)) == 1 and len(results[0]) == 64
    assert len(sheet_rows(path, "orders")) == 1
    assert len(sheet_rows(path, "items")) == 2
