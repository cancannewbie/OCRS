"""Synthetic integration tests for durable review, concurrency and recovery."""

import copy
import io
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Barrier
from typing import Any

import pytest
from openpyxl import load_workbook
from PIL import Image
from pydantic import ValidationError

from ocrs.config import Settings
from ocrs.domain import ConfirmationError
from ocrs.exporter import ExportError
from ocrs.providers import ProviderError
from ocrs.service import AppError, Confirmation, Service
from ocrs.storage import connect, migrate, transaction


@pytest.fixture
def service(tmp_path: Path) -> Service:
    migrate(tmp_path)
    return Service(Settings(data_dir=tmp_path, token="synthetic-local-test-token-" * 2))


def image_bytes() -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (12, 12), "white").save(stream, format="PNG")
    return stream.getvalue()


def ready(service: Service, label: str = "fictional-source") -> dict[str, Any]:
    task, _ = service.ingest(image_bytes(), "synthetic.png", label)
    assert service.process_one()
    task = service.task(task["id"])
    assert task["status"] == "review_required"
    return task


def confirmation(task: dict[str, Any], key: str = "request-synthetic-0001") -> Confirmation:
    return Confirmation(
        expected_version=task["version"],
        idempotency_key=key,
        actor="fictional-operator",
        reason="Reviewed a synthetic test image",
        events=task["candidate"]["events"],
    )


def counts(service: Service) -> tuple[int, int, int, int]:
    with connect(service.root) as db:
        return tuple(
            db.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
            for table in ("orders", "items", "events", "outbox")
        )


def test_human_review_is_required_before_export(service: Service) -> None:
    task = ready(service)
    assert service.orders() == []
    empty_export = service.export()
    assert empty_export["order_count"] == 0
    result = service.confirm(task["id"], confirmation(task))
    order = service.orders()[0]
    assert result == {"orders": [order["id"]]}
    assert order["version"] == 1
    assert order["items"][0]["quantity"] == "2"
    assert order["items"][0]["unit_price"] == "12.50"
    assert counts(service) == (1, 1, 1, 1)
    exported = service.export()
    assert exported["order_count"] == exported["item_count"] == 1
    book = load_workbook(service.root / "exports" / exported["filename"])
    try:
        assert book["orders"]["A2"].value == order["id"]
        assert book["orders"]["K2"].value == "25.00"
        assert book["items"]["B2"].value == order["items"][0]["line_id"]
    finally:
        book.close()
    assert service.status()["pending_export_events"] == 0


def test_upload_deduplication_is_scoped_by_source(service: Service) -> None:
    task, duplicate = service.ingest(image_bytes(), "../../fictional.png", "one")
    assert not duplicate
    repeated, duplicate = service.ingest(image_bytes(), "another.png", "one")
    assert duplicate and repeated["id"] == task["id"]
    different, duplicate = service.ingest(image_bytes(), "another.png", "two")
    assert not duplicate and different["id"] != task["id"]
    assert task["sources"][0]["filename"] == "fictional.png"
    assert len(list((service.root / "images").glob("*.png"))) == 2


def test_concurrent_uploads_create_one_source(service: Service) -> None:
    barrier = Barrier(2)

    def ingest() -> tuple[dict[str, Any], bool]:
        barrier.wait()
        return service.ingest(image_bytes(), "fictional.png", "same-source")

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: ingest(), range(2)))
    assert results[0][0]["id"] == results[1][0]["id"]
    assert sorted(duplicate for _, duplicate in results) == [False, True]
    assert len(service.tasks()) == 1


def test_confirmation_is_idempotent_but_changed_payload_conflicts(service: Service) -> None:
    task = ready(service)
    request = confirmation(task)
    result = service.confirm(task["id"], request)
    assert service.confirm(task["id"], request) == result
    changed = request.model_copy(update={"reason": "Different review reason"})
    with pytest.raises(AppError) as caught:
        service.confirm(task["id"], changed)
    assert caught.value.code == "IDEMPOTENCY_CONFLICT"
    assert counts(service) == (1, 1, 1, 1)


def test_concurrent_confirmation_allows_only_one_version(service: Service) -> None:
    task = ready(service)
    barrier = Barrier(2)

    def confirm(index: int) -> dict[str, Any] | str:
        request = confirmation(task, f"parallel-request-{index}")
        barrier.wait()
        try:
            return service.confirm(task["id"], request)
        except AppError as error:
            return error.code

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(confirm, range(2)))
    assert sum(isinstance(result, dict) for result in results) == 1
    assert "VERSION_CONFLICT" in results
    assert counts(service) == (1, 1, 1, 1)


def test_concurrent_identical_confirmation_reuses_response(service: Service) -> None:
    task = ready(service)
    request = confirmation(task)
    barrier = Barrier(2)

    def confirm() -> dict[str, Any]:
        barrier.wait()
        return service.confirm(task["id"], request)

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: confirm(), range(2)))
    assert results[0] == results[1]
    assert counts(service) == (1, 1, 1, 1)


def test_business_duplicate_requires_explicit_override(service: Service) -> None:
    original = ready(service, "one")
    service.confirm(original["id"], confirmation(original))
    duplicate = ready(service, "two")
    request = confirmation(duplicate, "duplicate-request-0002")
    with pytest.raises(AppError) as caught:
        service.confirm(duplicate["id"], request)
    assert caught.value.code == "DUPLICATE_WARNING"
    assert caught.value.details["order_ids"] == [service.orders()[0]["id"]]
    assert service.task(duplicate["id"])["status"] == "review_required"
    assert counts(service) == (1, 1, 1, 1)
    request.acknowledge_duplicates = True
    service.confirm(duplicate["id"], request)
    assert counts(service) == (2, 2, 2, 2)


def test_duplicate_events_in_same_confirmation_require_override(service: Service) -> None:
    task = ready(service)
    request = confirmation(task)
    request.events.append(request.events[0].model_copy(deep=True))
    with pytest.raises(AppError) as caught:
        service.confirm(task["id"], request)
    assert caught.value.code == "DUPLICATE_WARNING"
    assert counts(service) == (0, 0, 0, 0)
    request.acknowledge_duplicates = True
    service.confirm(task["id"], request)
    assert counts(service) == (2, 2, 2, 2)


def test_amend_cancel_keep_ids_versions_audit_and_outbox(service: Service) -> None:
    original = ready(service, "one")
    service.confirm(original["id"], confirmation(original))
    order = service.orders()[0]
    line_id = order["items"][0]["line_id"]
    task = ready(service, "amend")
    request = confirmation(task, "amend-request-0002")
    event = request.events[0]
    event.action = "amend"
    event.target_order_id = order["id"]
    event.expected_order_version = 1
    event.reason = "Correct fictional quantity"
    event.items[0].line_id = line_id
    event.items[0].quantity = "3.125"
    service.confirm(task["id"], request)
    amended = service.orders()[0]
    assert amended["version"] == 2
    assert amended["items"][0]["line_id"] == line_id
    assert amended["items"][0]["quantity"] == "3.125"
    cancellation = ready(service, "cancel")
    request = confirmation(cancellation, "cancel-request-0003")
    event = request.events[0]
    event.action = "cancel"
    event.target_order_id = order["id"]
    event.expected_order_version = 2
    event.reason = "Cancel fictional example"
    event.items = []
    service.confirm(cancellation["id"], request)
    cancelled = service.orders()[0]
    assert cancelled["status"] == "cancelled" and cancelled["version"] == 3
    assert cancelled["items"][0]["line_id"] == line_id
    assert counts(service) == (1, 1, 3, 3)
    with connect(service.root) as db:
        events = db.execute("SELECT * FROM events ORDER BY version").fetchall()
        assert [row["action"] for row in events] == ["create", "amend", "cancel"]
        assert [row["version"] for row in events] == [1, 2, 3]
        assert all(row["actor"] == "fictional-operator" for row in events)
        assert json.loads(events[0]["payload"])["items"][0]["quantity"] == "2"
    exported = service.export()
    book = load_workbook(service.root / "exports" / exported["filename"])
    try:
        assert book["orders"]["C2"].value == "cancelled"
        assert book["items"]["D2"].value == "cancelled"
    finally:
        book.close()


def test_invalid_second_action_rolls_back_entire_confirmation(service: Service) -> None:
    task = ready(service)
    request = confirmation(task)
    invalid = request.events[0].model_copy(deep=True)
    invalid.action = "cancel"
    invalid.target_order_id = "missing-order"
    invalid.expected_order_version = 1
    invalid.reason = "Fictional test cancellation"
    request.events.append(invalid)
    with pytest.raises(AppError) as caught:
        service.confirm(task["id"], request)
    assert caught.value.code == "VERSION_CONFLICT"
    assert counts(service) == (0, 0, 0, 0)
    assert service.task(task["id"])["status"] == "review_required"
    with connect(service.root) as db:
        assert db.execute("SELECT count(*) FROM requests").fetchone()[0] == 0


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("sku", None),
        ("sku", "UNKNOWN"),
        ("quantity", None),
        ("quantity", "0"),
        ("unit_price", None),
    ],
)
def test_missing_or_unknown_fields_never_create_order(
    service: Service,
    field: str,
    value: object,
) -> None:
    task = ready(service)
    request = confirmation(task)
    setattr(request.events[0].items[0], field, value)
    with pytest.raises(ConfirmationError):
        service.confirm(task["id"], request)
    assert counts(service) == (0, 0, 0, 0)


def test_binary_float_is_rejected_at_request_boundary(service: Service) -> None:
    task = ready(service)
    payload = confirmation(task).model_dump(mode="json")
    payload["events"][0]["items"][0]["unit_price"] = 0.1
    with pytest.raises(ValidationError):
        Confirmation.model_validate(payload)


def test_evidence_from_another_source_is_rejected(service: Service) -> None:
    task = ready(service)
    request = confirmation(task)
    request.events[0].evidence[0].source_id = "unrelated-source"
    with pytest.raises(AppError) as caught:
        service.confirm(task["id"], request)
    assert caught.value.code == "EVIDENCE_INVALID"
    assert counts(service) == (0, 0, 0, 0)


def test_rejected_task_cannot_be_confirmed_or_retried(service: Service) -> None:
    task = ready(service)
    rejected = service.reject(task["id"], task["version"], "Synthetic rejection")
    assert rejected["status"] == "rejected"
    with pytest.raises(AppError, match="任务已被处理"):
        service.confirm(task["id"], confirmation(task))
    with pytest.raises(AppError) as caught:
        service.retry(task["id"])
    assert caught.value.code == "STATE_CONFLICT"
    assert not service.process_one()
    assert counts(service) == (0, 0, 0, 0)
    with connect(service.root) as db:
        assert db.execute("SELECT kind FROM audit").fetchone()[0] == "rejected"


def test_recovery_marks_interrupted_work_and_retains_confirmed_order(service: Service) -> None:
    task = ready(service)
    service.confirm(task["id"], confirmation(task))
    interrupted, _ = service.ingest(image_bytes(), "sample.png", "interrupted")
    with transaction(service.root) as db:
        db.execute("UPDATE tasks SET status='recognizing' WHERE id=?", (interrupted["id"],))
        db.execute(
            "INSERT INTO exports(id,status,path,created_at) VALUES('crash-export',"
            "'writing','exports/crash.xlsx','2000-01-01T00:00:00+00:00')"
        )
        db.execute("UPDATE outbox SET status='writing',export_id='crash-export'")
    recovered = Service(service.settings)
    recovered.recover()
    failed = recovered.task(interrupted["id"])
    assert failed["status"] == "failed" and failed["error_code"] == "INTERRUPTED"
    assert counts(service) == (1, 1, 1, 1)
    assert service.task(task["id"])["status"] == "confirmed"
    with connect(service.root) as db:
        assert db.execute("SELECT status FROM exports").fetchone()[0] == "failed"
        assert db.execute("SELECT status,export_id FROM outbox").fetchone()[:] == ("pending", None)
    recovered.retry(interrupted["id"])
    assert recovered.process_one()
    assert recovered.task(interrupted["id"])["status"] == "review_required"


def test_provider_failure_is_safe_and_manual_retries_are_bounded(
    service: Service,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    task, _ = service.ingest(image_bytes(), "fictional.png", "failure")

    def fail(*args: object) -> Any:
        raise ProviderError("provider_timeout")

    monkeypatch.setattr(service.provider, "recognize", fail)
    for attempt in range(1, 6):
        assert service.process_one()
        current = service.task(task["id"])
        assert current["status"] == "failed" and current["attempts"] == attempt
        assert current["error_code"] == "provider_timeout"
        assert not service.process_one()
        if attempt < 5:
            service.retry(task["id"])
    with pytest.raises(AppError) as caught:
        service.retry(task["id"])
    assert caught.value.code == "RETRY_LIMIT"
    assert service.settings.token not in caplog.text
    assert counts(service) == (0, 0, 0, 0)


def test_unexpected_provider_exception_never_exposes_raw_body(
    service: Service,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    task, _ = service.ingest(image_bytes(), "fictional.png", "private-error")

    def fail(*args: object) -> Any:
        raise RuntimeError("SYNTHETIC_SECRET_AND_CUSTOMER_TEXT")

    monkeypatch.setattr(service.provider, "recognize", fail)
    service.process_one()
    assert service.task(task["id"])["error_code"] == "RECOGNITION_INTERNAL"
    assert "SYNTHETIC_SECRET_AND_CUSTOMER_TEXT" not in caplog.text


def test_export_failure_preserves_order_and_retries_outbox(
    service: Service,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import ocrs.service as service_module

    task = ready(service)
    service.confirm(task["id"], confirmation(task))
    snapshot = copy.deepcopy(service.orders())
    writer = service_module.write_workbook

    def fail(*args: object) -> Any:
        raise ExportError("export_locked", "Synthetic file lock", retryable=True)

    monkeypatch.setattr(service_module, "write_workbook", fail)
    with pytest.raises(AppError) as caught:
        service.export()
    assert caught.value.code == "EXPORT_FAILED"
    assert service.orders() == snapshot
    with connect(service.root) as db:
        assert db.execute("SELECT status FROM outbox").fetchone()[0] == "failed"
    monkeypatch.setattr(service_module, "write_workbook", writer)
    assert service.export()["status"] == "completed"
    assert service.status()["pending_export_events"] == 0
    assert counts(service) == (1, 1, 1, 1)


def test_stale_order_version_and_foreign_line_id_leave_original_unchanged(service: Service) -> None:
    original = ready(service, "original")
    service.confirm(original["id"], confirmation(original))
    snapshot = copy.deepcopy(service.orders())
    order = snapshot[0]
    task = ready(service, "amend")
    request = confirmation(task, "stale-amend-request")
    event = request.events[0]
    event.action = "amend"
    event.target_order_id = order["id"]
    event.expected_order_version = 2
    event.reason = "Correct synthetic order"
    with pytest.raises(AppError) as caught:
        service.confirm(task["id"], request)
    assert caught.value.code == "VERSION_CONFLICT"
    event.expected_order_version = 1
    event.items[0].line_id = "foreign-line-id"
    with pytest.raises(AppError) as caught:
        service.confirm(task["id"], request)
    assert caught.value.code == "LINE_ID_INVALID"
    assert service.orders() == snapshot
    assert counts(service) == (1, 1, 1, 1)
    assert service.task(task["id"])["status"] == "review_required"


@pytest.mark.parametrize("action", ["amend", "cancel"])
def test_event_and_outbox_failure_rolls_back_order_change(
    service: Service,
    action: str,
) -> None:
    import sqlite3

    original = ready(service, "original")
    service.confirm(original["id"], confirmation(original))
    snapshot = copy.deepcopy(service.orders())
    task = ready(service, f"change-{action}")
    request = confirmation(task, f"outbox-failure-{action}")
    event = request.events[0]
    event.action = action
    event.target_order_id = snapshot[0]["id"]
    event.expected_order_version = 1
    event.reason = "Fictional change for atomicity test"
    if action == "amend":
        event.items[0].line_id = snapshot[0]["items"][0]["line_id"]
        event.items[0].quantity = "4"
    with transaction(service.root) as db:
        db.execute(
            "CREATE TRIGGER simulate_outbox_failure BEFORE INSERT ON outbox "
            "BEGIN SELECT RAISE(ABORT, 'synthetic outbox failure'); END"
        )
    with pytest.raises(sqlite3.IntegrityError):
        service.confirm(task["id"], request)
    assert service.orders() == snapshot
    assert counts(service) == (1, 1, 1, 1)
    assert service.task(task["id"])["status"] == "review_required"
    with transaction(service.root) as db:
        assert db.execute("SELECT count(*) FROM requests").fetchone()[0] == 1
        db.execute("DROP TRIGGER simulate_outbox_failure")
    service.confirm(task["id"], request)
    assert service.orders()[0]["version"] == 2
    assert counts(service) == (1, 1, 2, 2)


def test_export_crash_after_publish_is_reconciled_from_database(service: Service) -> None:
    task = ready(service)
    service.confirm(task["id"], confirmation(task))
    exported = service.export()
    first_path = service.root / "exports" / exported["filename"]
    published = first_path.read_bytes()
    with transaction(service.root) as db:
        db.execute("UPDATE exports SET status='writing' WHERE id=?", (exported["id"],))
        db.execute("UPDATE outbox SET status='writing'")
    recovered = Service(service.settings)
    recovered.recover()
    rebuilt = recovered.export()
    assert rebuilt["order_count"] == rebuilt["item_count"] == 1
    assert first_path.read_bytes() == published
    assert counts(service) == (1, 1, 1, 1)
    assert recovered.status()["pending_export_events"] == 0


@pytest.mark.parametrize("field,value", [("actor", "   "), ("reason", "\t")])
def test_blank_audit_identity_or_reason_rejected(service: Service, field: str, value: str) -> None:
    task = ready(service)
    request = confirmation(task)
    setattr(request, field, value)
    with pytest.raises(AppError, match="不得为空白"):
        service.confirm(task["id"], request)
    assert counts(service) == (0, 0, 0, 0)


def test_evidence_digest_changes_fail_before_provider(service: Service) -> None:
    task, _ = service.ingest(image_bytes(), "synthetic.png", "digest-check")
    with connect(service.root) as db:
        path = db.execute("SELECT path FROM sources WHERE id=?", (task["source_id"],)).fetchone()[0]
    (service.root / path).write_bytes(b"modified-evidence")
    service.process_one()
    assert service.task(task["id"])["error_code"] == "EVIDENCE_CHANGED"


def test_same_batch_external_id_duplicate_even_when_other_fields_differ(service: Service) -> None:
    task = ready(service)
    request = confirmation(task)
    request.events[0].external_id = "fictional-reference-001"
    other = request.events[0].model_copy(deep=True)
    other.customer = "Different fictional customer"
    request.events.append(other)
    with pytest.raises(AppError) as error:
        service.confirm(task["id"], request)
    assert error.value.code == "DUPLICATE_WARNING"
    assert counts(service) == (0, 0, 0, 0)
